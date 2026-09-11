"""Tests for yotext.docx."""

import pytest

pytest.importorskip("docx")

from docx import Document

from yotext import standardize, strip_diacritics
from yotext.docx import read_docx, process_docx

MESSY = "e\u0300\u0329ko\u0301\u0331"
CANONICAL = "\u1eb9\u0300k\u1ecd\u0301"
ITALIC_WORD = "\u1ecd\u0300w\u1ecd\u0300"


def _make(path, paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))
    return path


def test_read_docx_joins_paragraphs_with_newlines(tmp_path):
    path = _make(tmp_path / "in.docx", ["first", "second", "third"])

    assert read_docx(path) == "first\nsecond\nthird"


def test_process_docx_standardizes(tmp_path):
    src = _make(tmp_path / "in.docx", [MESSY])
    out = tmp_path / "out.docx"

    process_docx(src, out, standardize)

    assert read_docx(out) == CANONICAL


def test_process_docx_strip_diacritics(tmp_path):
    src = _make(tmp_path / "in.docx", [CANONICAL])
    out = tmp_path / "out.docx"

    process_docx(src, out, strip_diacritics)

    assert read_docx(out) == "eko"


def test_run_formatting_survives(tmp_path):
    src = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    document = Document()
    paragraph = document.add_paragraph()
    bold_run = paragraph.add_run(MESSY)
    bold_run.bold = True
    paragraph.add_run(" ")
    italic_run = paragraph.add_run(ITALIC_WORD)
    italic_run.italic = True
    document.save(str(src))

    process_docx(src, out, standardize)

    runs = Document(str(out)).paragraphs[0].runs
    assert len(runs) == 3
    assert runs[0].bold is True
    assert runs[0].italic is not True
    assert runs[0].text == CANONICAL
    assert runs[1].text == " "
    assert runs[2].italic is True
    assert runs[2].bold is not True
    assert runs[2].text == ITALIC_WORD


def test_missing_docx_raises_file_not_found(tmp_path):
    missing = tmp_path / "nope.docx"

    with pytest.raises(FileNotFoundError) as exc_info:
        read_docx(missing)
    assert "nope.docx" in str(exc_info.value)

    with pytest.raises(FileNotFoundError):
        process_docx(missing, tmp_path / "out.docx", standardize)
