"""Read and rewrite .docx files while preserving their formatting."""

import errno
from pathlib import Path

_INSTALL_HINT = (
    "Working with .docx files needs python-docx, which is not installed. "
    "Install it with: pip install yotext[docx]"
)


def _open_document(path):
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError(_INSTALL_HINT) from exc
    # python-docx reports a missing file as its own PackageNotFoundError, which
    # callers would have to import python-docx to catch. Check first so a
    # missing file is always a plain FileNotFoundError carrying the path.
    if not Path(path).is_file():
        raise FileNotFoundError(errno.ENOENT, "No such file or directory", str(path))
    return Document(str(path))


def read_docx(path) -> str:
    """Return the text of every paragraph in the document, joined by newlines."""
    document = _open_document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def process_docx(path, out_path, func) -> None:
    """Apply func to the text of each run and save the document to out_path.

    Each run is rewritten in place rather than the paragraph as a whole.
    Assigning to paragraph.text would collapse the paragraph into a single
    run and discard the character formatting of everything inside it.
    """
    document = _open_document(path)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.text = func(run.text)
    document.save(str(out_path))
